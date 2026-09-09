"""Lead upload parsing and the validation pipeline (MVP section 3).

Pipeline order follows the MVP:
    validate file -> validate phone -> normalize +91 -> duplicate check ->
    previous-lead check -> consent check -> DNC/opt-out check -> READY
"""

import csv
import io

from sqlalchemy import select
from sqlalchemy.orm import Session

from backend.auth.models import User
from backend.compliance.service import is_suppressed, log_action
from backend.customers.models import Customer
from backend.leads.models import (
    Lead,
    LeadRejection,
    LeadStatus,
    LeadUpload,
    RejectionReason,
    ServiceType,
    UploadStatus,
)
from backend.leads.phone import normalize_indian_mobile

REQUIRED_COLUMNS = ["customer_name", "mobile_number", "lead_source", "consent_status"]
OPTIONAL_COLUMNS = [
    "lead_id",
    "alternate_number",
    "interested_service",
    "language",
    "city",
    "district",
    "mandal",
    "pincode",
    "monthly_bill",
    "monthly_units",
    "campaign",
    "remarks",
    "consent_date",
    "consent_source",
]
TEMPLATE_COLUMNS = [
    "Customer_Name",
    "Mobile_Number",
    "Lead_Source",
    "Consent_Status",
    "Lead_ID",
    "Alternate_Number",
    "Interested_Service",
    "Language",
    "City",
    "District",
    "Mandal",
    "Pincode",
    "Monthly_Bill",
    "Monthly_Units",
    "Campaign",
    "Remarks",
    "Consent_Date",
    "Consent_Source",
]

_AFFIRMATIVE_CONSENT = {"yes", "y", "true", "1", "opted_in", "opted-in", "optedin", "given", "consented"}

# Lead statuses that make a new upload of the same phone a duplicate rather
# than a fresh lead.
_OPEN_LEAD_STATUSES = (LeadStatus.READY, LeadStatus.IN_CAMPAIGN)


class FileFormatError(Exception):
    pass


def _normalize_header(header: str) -> str:
    return str(header).strip().lower().replace(" ", "_").replace("-", "_")


def _clean(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _to_number(value) -> float | None:
    text = _clean(value)
    if text is None:
        return None
    try:
        return float(text.replace(",", ""))
    except ValueError:
        return None


def _to_service(value) -> ServiceType | None:
    text = _clean(value)
    if text is None:
        return None
    key = text.upper().replace(" ", "_").replace("-", "_")
    try:
        return ServiceType(key)
    except ValueError:
        return None


def parse_file(filename: str, content: bytes) -> list[dict]:
    """Parse a CSV or XLSX file into a list of dicts with normalized keys."""
    name = filename.lower()
    if name.endswith(".csv"):
        rows = _parse_csv(content)
    elif name.endswith(".xlsx"):
        rows = _parse_xlsx(content)
    else:
        raise FileFormatError("Unsupported file type — upload a .csv or .xlsx file")

    if not rows:
        raise FileFormatError("The file contains no data rows")

    headers = set(rows[0].keys())
    missing = [c for c in REQUIRED_COLUMNS if c not in headers]
    if missing:
        raise FileFormatError(
            "Missing required columns: " + ", ".join(sorted(missing))
        )
    return rows


def _parse_csv(content: bytes) -> list[dict]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise FileFormatError("CSV file must be UTF-8 encoded")
    reader = csv.DictReader(io.StringIO(text))
    if not reader.fieldnames:
        raise FileFormatError("The file has no header row")
    return [
        {_normalize_header(k): v for k, v in row.items() if k is not None}
        for row in reader
    ]


def _parse_xlsx(content: bytes) -> list[dict]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise FileFormatError("Could not read the Excel file")
    sheet = workbook.active
    rows_iter = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        raise FileFormatError("The file has no header row")
    headers = [_normalize_header(h) if h is not None else "" for h in header_row]
    rows = []
    for values in rows_iter:
        if values is None or all(v is None or str(v).strip() == "" for v in values):
            continue
        rows.append({h: v for h, v in zip(headers, values) if h})
    return rows


def import_leads(db: Session, *, user: User, filename: str, content: bytes) -> LeadUpload:
    """Run the full validation pipeline over an uploaded file.

    Commits once at the end; the whole upload succeeds or fails together.
    """
    upload = LeadUpload(filename=filename, uploaded_by_id=user.id)
    db.add(upload)

    try:
        rows = parse_file(filename, content)
    except FileFormatError as exc:
        upload.status = UploadStatus.FAILED
        upload.error = str(exc)
        log_action(
            db, actor=user, action="LEAD_UPLOAD_FAILED",
            entity=f"upload:{filename}", details={"error": str(exc)},
        )
        db.commit()
        db.refresh(upload)
        return upload

    db.flush()  # assign upload.id for FKs
    upload.total_rows = len(rows)
    seen_in_file: set[str] = set()

    for index, row in enumerate(rows, start=2):  # row 1 is the header
        name = _clean(row.get("customer_name"))
        raw_phone = _clean(row.get("mobile_number"))
        source = _clean(row.get("lead_source"))
        consent = _clean(row.get("consent_status"))

        def reject(reason: RejectionReason, details: str, counter: str):
            db.add(
                LeadRejection(
                    upload_id=upload.id,
                    row_number=index,
                    customer_name=name,
                    mobile_number=raw_phone,
                    reason=reason,
                    details=details,
                )
            )
            setattr(upload, counter, getattr(upload, counter) + 1)

        # 1. Required fields
        missing = [
            label
            for label, value in (
                ("Customer_Name", name),
                ("Mobile_Number", raw_phone),
                ("Lead_Source", source),
                ("Consent_Status", consent),
            )
            if value is None
        ]
        if missing:
            reject(
                RejectionReason.MISSING_REQUIRED_FIELD,
                "Missing: " + ", ".join(missing),
                "missing_field_count",
            )
            continue

        # 2-3. Validate and normalize phone
        phone = normalize_indian_mobile(raw_phone)
        if phone is None:
            reject(
                RejectionReason.INVALID_PHONE,
                "Not a valid Indian mobile number",
                "invalid_phone_count",
            )
            continue

        # 4. Duplicate within this file
        if phone in seen_in_file:
            reject(
                RejectionReason.DUPLICATE_IN_FILE,
                "Number appears earlier in this file",
                "duplicate_count",
            )
            continue
        seen_in_file.add(phone)

        # 5. Previous-lead check: an open lead already exists for this customer
        customer = db.scalar(select(Customer).where(Customer.phone == phone))
        if customer is not None:
            open_lead = db.scalar(
                select(Lead)
                .where(Lead.customer_id == customer.id, Lead.status.in_(_OPEN_LEAD_STATUSES))
                .limit(1)
            )
            if open_lead is not None:
                reject(
                    RejectionReason.DUPLICATE_EXISTING,
                    f"Open lead {open_lead.lead_ref} already exists for this customer",
                    "duplicate_count",
                )
                continue

        # 6. Consent check
        if consent.lower() not in _AFFIRMATIVE_CONSENT:
            reject(
                RejectionReason.CONSENT_MISSING,
                f"Consent_Status '{consent}' is not an affirmative consent",
                "consent_issue_count",
            )
            continue

        # 7. DNC / opt-out suppression check
        if (customer is not None and customer.opted_out) or is_suppressed(db, phone):
            reject(
                RejectionReason.OPTED_OUT,
                "Number is on the do-not-call suppression list",
                "opted_out_count",
            )
            continue

        # Ready for campaign — attach to customer history (create if new)
        if customer is None:
            customer = Customer(phone=phone, name=name)
            db.add(customer)
            db.flush()
        elif customer.name != name:
            customer.name = name

        alternate = _clean(row.get("alternate_number"))
        lead = Lead(
            customer_id=customer.id,
            upload_id=upload.id,
            source=source,
            consent_status=consent,
            consent_date=_clean(row.get("consent_date")),
            consent_source=_clean(row.get("consent_source")),
            interested_service=_to_service(row.get("interested_service")),
            language=_clean(row.get("language")),
            alternate_number=normalize_indian_mobile(alternate) or alternate,
            city=_clean(row.get("city")),
            district=_clean(row.get("district")),
            mandal=_clean(row.get("mandal")),
            pincode=_clean(row.get("pincode")),
            monthly_bill=_to_number(row.get("monthly_bill")),
            monthly_units=_to_number(row.get("monthly_units")),
            campaign_hint=_clean(row.get("campaign")),
            remarks=_clean(row.get("remarks")),
            status=LeadStatus.READY,
        )
        db.add(lead)
        db.flush()
        lead.lead_ref = f"SW-{10000 + lead.id}"
        upload.valid_count += 1

    log_action(
        db,
        actor=user,
        action="LEAD_UPLOAD_COMPLETED",
        entity=f"upload:{upload.id}",
        details={
            "filename": filename,
            "total": upload.total_rows,
            "valid": upload.valid_count,
            "duplicate": upload.duplicate_count,
            "invalid_phone": upload.invalid_phone_count,
            "opted_out": upload.opted_out_count,
            "consent_issue": upload.consent_issue_count,
            "missing_field": upload.missing_field_count,
        },
    )
    db.commit()
    db.refresh(upload)
    return upload
